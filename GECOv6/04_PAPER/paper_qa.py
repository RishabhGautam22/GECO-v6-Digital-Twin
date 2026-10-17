"""Validate final paper exports against frozen aggregates and record visual QA."""
from pathlib import Path
import hashlib, json, re, sys, subprocess, platform
from zipfile import ZipFile
from datetime import datetime, timezone
from lxml import etree
from docx import Document
import docx, pypdf
from PIL import Image

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'04_PAPER'; WORK=OUT/'paper_work'
QA=ROOT/'08_REVIEW_AND_QA/paper'; QA.mkdir(parents=True,exist_ok=True)
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def read(p):return json.loads((ROOT/p).read_text())
def norm(s):return re.sub(r'\s+','',s).replace('−','-').replace('–','-')
pdf=OUT/'GECOv6_Manuscript.pdf';word=OUT/'GECOv6_Manuscript.docx';md=OUT/'GECOv6_MANUSCRIPT.md'
doc=Document(word);reader=pypdf.PdfReader(pdf)
pdftext='\n'.join(p.extract_text() for p in reader.pages)
with ZipFile(word) as z:
    x=etree.fromstring(z.read('word/document.xml'));ns=x.nsmap
    wordtext=' '.join(x.xpath('.//w:t/text()',namespaces=ns))
    rel=etree.fromstring(z.read('word/_rels/document.xml.rels'))
    urls=[a.get('Target') for a in rel if a.get('TargetMode')=='External']
assert len(reader.pages)==10
assert len(doc.tables)==4
assert len(x.xpath('.//w:drawing',namespaces=ns))==3
assert len(x.xpath('.//m:oMath',namespaces=ns))==73
labels=[]
for p in x.xpath('.//w:p',namespaces=ns):
    text=''.join(p.xpath('.//w:t/text()',namespaces=ns)).strip()
    if text in [f'({i})' for i in range(1,8)]:
        assert p.xpath('./m:oMath',namespaces=ns)
        labels.append(text)
assert labels==[f'({i})' for i in range(1,8)]
assert all('\\' not in u for u in urls)
assert any('JCGM_100_2008_E.pdf' in u for u in urls)
assert all(f'M{i:02}' in wordtext and f'M{i:02}' in pdftext for i in range(1,17))
assert len(doc.tables[3].rows)==7
assert all('Pending' in ' '.join(c.text for c in r.cells) for r in doc.tables[3].rows[1:])
assert wordtext.count('Fig. 1.')==1 and wordtext.count('Fig. 2.')==1
assert wordtext.count('Fig. 3.')==1 and 'Figure 3 links' in wordtext
assert len(x.xpath('.//w:pPr/w:shd[@w:fill="FFF2CC"]',namespaces=ns))==2

keys=re.findall(r'\\bibcite\{(R\d+)\}\{(\d+)\}',(WORK/'GECOv6_Manuscript.aux').read_text())
refmap={k:int(v) for k,v in keys}
bookmarks=x.xpath('.//w:bookmarkStart/@w:name',namespaces=ns)
wordkeys=[s[4:] for s in bookmarks if s.startswith('ref-R')]
assert wordkeys==[k for k,v in sorted(refmap.items(),key=lambda a:a[1])]
assert set(refmap)=={f'R{i:02}' for i in range(1,16)}
assert not re.search(r'Overfull|undefined citations|Citation .* undefined', (WORK/'GECOv6_Manuscript.log').read_text(),re.I)

agg=read('03_EXPERIMENTS/filter_calibration/AGGREGATE_RESULTS.json')
main=read('03_EXPERIMENTS/comparative_study/RESULT_SUMMARY.json')
supp=read('03_EXPERIMENTS/supplement_fasting/RESULT_SUMMARY.json')
assert agg['n_evaluated']==31 and agg['all_test_pairs']==77110
assert agg['selection_counts']=={'trend':30,'persistence':1}
assert main['completed_runs']==5184 and main['main_runs']==4608 and main['feedback_stop_runs']==576 and main['failures']==0
assert main['main_runs_with_invalid_physiology']==2108
assert supp['completed_runs']==2048 and supp['failed_runs']==0 and supp['runs_domain_invalid']==0 and supp['runs_negative_egp']==0
for n in ['77,110','4,608','576','5,184','2,985,984','2,108','138,103','1,327,104','579','133,394','34,731','2,048','1,179,648','18,648','74.03','203.03']:
    assert norm(n) in norm(pdftext) and norm(n) in norm(wordtext),n
medkeys=['mechanistic_rmse','persistence_rmse','trend_rmse','selected_rmse']
for row,key in zip(doc.tables[1].rows[1:],medkeys):
    assert row.cells[1].text==f"{agg['participant_median'][key]:.2f}"
for i,k in [(1,'mechanistic'),(3,'trend'),(4,'selected')]:
    a=agg['paired_rmse_comparisons'][k+'_minus_persistence']
    expected=f"{a['mean_difference_mgdl']:+.2f} [{a['participant_bootstrap95'][0]:.2f},{a['participant_bootstrap95'][1]:.2f}]"
    assert norm(doc.tables[1].rows[i].cells[2].text)==norm(expected)

mainpairs=read('03_EXPERIMENTS/comparative_study/PAIRED_PARAMETER_UNIT_CI.json')
condition={'Clean':'clean','2× actuator input':'excess_2x','Occlusion':'occlusion','Stuck-on addition':'stuck_on','CGM +40':'cgm_bias_40','CGM missing':'cgm_missing','2× input + CGM +25':'excess_bias_25','2× input + bias + report replacement':'excess_bias_report_mask'}
cells=[]
for row in doc.tables[2].rows[1:]:
    family='geco_original' if row.cells[0].text=='Original' else 'hovorka_2004'
    cond=condition[row.cells[1].text]
    for col,metric in [(2,'stepwise_excess_vs_clean_u'),(3,'stepwise_deficit_vs_clean_u')]:
        a=next(q for q in mainpairs if q['mode']=='main' and q['family']==family and q['condition']==cond and q['arm_a']=='full_supervised' and q['arm_b']=='simple_supervised' and q['metric']==metric)
        expected=f"{a['mean_paired_difference']:+.3f} [{a['ci95_low']:+.3f},{a['ci95_high']:+.3f}]"
        assert row.cells[col].text==expected
        assert norm(expected) in norm(pdftext)
        assert a['n_parameter_units']==24 and a['paired_seed_runs']==48
        cells.append({'family':family,'condition':cond,'metric':metric,'printed':expected,'matched':True})
assert len(cells)==32
supairs=read('03_EXPERIMENTS/supplement_fasting/PAIRED_PARAMETER_UNIT_CI.json')
suchecks=[]
for fam in ['geco_original','hovorka_2004']:
    for cond,metric in [('excess_bias_25','stepwise_excess_vs_clean_u'),('cgm_bias_40','stepwise_deficit_vs_clean_u')]:
        a=next(q for q in supairs if q['family']==fam and q['condition']==cond and q['metric']==metric and q['arm_a']=='full_supervised' and q['arm_b']=='masscusum_supervised')
        for field in ['mean_paired_difference','ci95_low','ci95_high']:
            value=f"{a[field]:.5f}"
            assert norm(value) in norm(wordtext) and norm(value) in norm(pdftext)
        suchecks.append({k:a[k] for k in ['family','condition','metric','mean_paired_difference','ci95_low','ci95_high']})

pdfnotes=[
 'Title, author placeholder, abstract numbers and scope legible.',
 'Prior-work limits, evidence table and native model equation (1) legible.',
 'Equations (2)–(5), calibration and information-boundary text legible.',
 'Pseudocode, grouped equation (6), superscripts, timing and equation (7) legible.',
 'Main invalid-domain findings and negative comparison remain explicit.',
 'Intentional float page: Table II, Fig. 1 and all 16 Table III rows fit; no overflow.',
 'Fig. 2, sensitivity and exploratory supplement legible; adverse costs retained.',
 'Both amber notices are intact, with all M01–M16 fields; Fig. 3 cross-reference and metrology contract legible.',
 'Full-width Fig. 3 conceptual rig and vector chain, six-row amber pending table and conclusion legible.',
 'Team disclosure placeholders and all 15 references legible; normal unused final-column space.'
]
wordnotes=[
 'Title, author placeholder, abstract and opening scope legible.',
 'Prior-work access limits, novelty boundaries and evidence definitions legible.',
 'Editable Table I and fault/controller definitions legible.',
 'Editable equations (1)–(3), timing, calibration and causal-input limits legible.',
 'Editable equations (4)–(5) and shared-authority rules legible.',
 'Pseudocode, grouped equation (6), sensitivities and simulation design legible.',
 'Equation (7), retained domain limits and cohort accounting legible; table kept together on next page.',
 'Whole Table II and Fig. 1 with explicit caption legible.',
 'Fig. 2 with explicit caption and main model-domain failures legible.',
 'Whole editable Table III with all 16 conditions; paired values and discussion legible.',
 'Sensitivity, exploratory supplement and complete first amber notice legible.',
 'Fig. 3 and caption are legible; conceptual rig and pending reconstruction stay explicit beside protocol prose.',
 'Complete M01–M16 amber notice and six-row amber table legible; discussion follows.',
 'Discussion, conclusion, provenance limits, human/team approval placeholders and first references legible.',
 'Remaining numbered references legible; corrected BIPM URL and normal unused final-page space.'
]
pages=[]
for kind,pattern,notes in [('IEEE PDF',WORK/'paper-final-*.png',pdfnotes),('DOCX rendered',WORK/'docx_final/page-*.png',wordnotes)]:
    pp=sorted(pattern.parent.glob(pattern.name),key=lambda p:int(re.search(r'(\d+)\.png$',p.name).group(1)))
    assert len(pp)==len(notes)
    for p,note in zip(pp,notes):
        with Image.open(p) as im:dims=list(im.size)
        pages.append({'format':kind,'page':int(re.search(r'(\d+)\.png$',p.name).group(1)),'path':str(p.relative_to(ROOT)),'sha256':sha(p),'pixels':dims,'viewed_individually':True,'inspection_tool':'view_image detail=original, forwarded at original detail','finding':note,'material_defect':False})
sources=[
 '03_EXPERIMENTS/filter_calibration/AGGREGATE_RESULTS.json',
 '03_EXPERIMENTS/filter_calibration/FROZEN_CALIBRATION.json',
 '03_EXPERIMENTS/comparative_study/RESULT_SUMMARY.json',
 '03_EXPERIMENTS/comparative_study/PAIRED_PARAMETER_UNIT_CI.json',
 '03_EXPERIMENTS/supplement_fasting/RESULT_SUMMARY.json',
 '03_EXPERIMENTS/supplement_fasting/PAIRED_PARAMETER_UNIT_CI.json',
 '08_REVIEW_AND_QA/filter/INDEPENDENT_SUPPLEMENT_REVIEW.md',
 '01_RESEARCH_PROTOCOL/literature_editorial/references.bib',
 '01_RESEARCH_PROTOCOL/literature_editorial/SOURCE_LEDGER.json',
 '06_RIG_AND_MEDIA/stills/rig_overview.png',
 '04_PAPER/paper_figures/Figure3_SOURCE.json',
 '04_PAPER/paper_figures/Figure3.pdf',
 '04_PAPER/paper_figures/Figure3.svg'
]
before=read('08_REVIEW_AND_QA/paper/PAPER_QA_BEFORE_FIG3.json')
prior_hash={p['path']:p['sha256'] for p in before['pages']}
unchanged=[{'format':p['format'],'page':p['page']} for p in pages if prior_hash.get(p['path'])==p['sha256']]
changed=[{'format':p['format'],'page':p['page']} for p in pages if prior_hash.get(p['path'])!=p['sha256']]
assert len(unchanged)==18 and len(changed)==7
receipt={
 'created_utc':datetime.now(timezone.utc).isoformat(),'status':'Final paper export QA passed; team/submission and physical work remain pending',
 'outputs':[{ 'path':str(p.relative_to(ROOT)),'sha256':sha(p),'bytes':p.stat().st_size} for p in [md,pdf,word,OUT/'GECOv6_Manuscript.tex',OUT/'paper_references.bib',OUT/'paper_build.py']],
 'page_counts':{'ieee_pdf':10,'word_rendered':15},'native_word':{'OMML_expressions':73,'numbered_display_equations':labels,'editable_tables':4,'figures':3,'explicit_figure_captions':3,'amber_pending_paragraphs':2,'pending_fields':16,'pending_endpoint_rows':6},
 'figure3_revision':{'unchanged_page_image_identities':unchanged,'changed_or_added_pages_reinspected':changed,'source_image_used_whole':True,'diagram_pdf_is_vector':True,'editable_svg_retained':True,'physical_validation_pending':True},
 'reference_key_to_ieee_number':refmap,'word_and_pdf_reference_order_match':True,'tableIII_all_32_numeric_cells':cells,'supplement_four_effects_and_intervals':suchecks,
 'checked_headlines':{'participants':31,'forecast_pairs':77110,'primary_total':5184,'primary_main':4608,'sensitivity':576,'main_domain_invalid_runs':2108,'supplement_runs':2048,'supplement_domain_invalid_runs':0},
 'sources':[{ 'path':s,'sha256':sha(ROOT/s)} for s in sources],
 'runtime':{'python':sys.executable,'python_version':platform.python_version(),'python_docx_version':docx.__version__,'python_docx_module':docx.__file__,'pandoc':'/opt/homebrew/bin/pandoc','pdf_engine':'/Library/TeX/texbin/pdflatex','word_renderer':'bundled documents render_docx.py and bundled soffice override','bundled_soffice':'/Users/rgpro/.cache/codex-runtimes/codex-primary-runtime/dependencies/bin/override/soffice','rasterizer':'/opt/homebrew/bin/pdftoppm','model_id':'unavailable in artifact tool results; not inferred','temperature':'unavailable','request_id':'unavailable'},
 'marker_metadata':{'docx':'Successful create marker exactly once before first authoring; no repeat','pdf':'Coordinator already invoked the PDF marker; this lane did not repeat it'},
 'closed_export_defects':['Raw variable and sensitivity powers converted to native mathematics.','Word equation numbers restored with native editable OMML.','Hovorka glucose equation grouped in one PDF column.','PDF figure/table overfull box removed without shrinking text.','First PDF amber pending notice kept intact.','Word figures given explicit captions, with limitations adjacent.','Word tables kept together instead of one-row carryovers.','Printed bibliography title capitalization, Dalla Man surname parsing, audit notes and BIPM URL escape corrected.'],
 'limitations':['Word layout was rendered with bundled LibreOffice; Microsoft Word was not launched in this lane. Native OMML and editable table structures were inspected.','IEEEtran style is a working journal manuscript format; a venue has not been selected and author/funding/ethics/AI-human-approval fields remain pending.','The review establishes export fidelity and bounded scientific reporting, not physiological validation, clinical benefit, physical rig performance, institutional endorsement, publication acceptance or originality certification.','Prior local exact-text overlap check covered 75 available texts on an earlier pre-supplement version; it was not proprietary-service, semantic or exhaustive plagiarism screening.','No external uploading, publication or third-party communication was performed.','After adding Fig. 3, PDF pages 1–7 and Word pages 1–11 are pixel-identical to their individually inspected prior versions; all seven changed/added pages were individually re-inspected.'],
 'pages':pages
}
(QA/'FINAL_PAPER_QA.json').write_text(json.dumps(receipt,indent=2)+'\n')
rows='\n'.join(f"| {p['format']} | {p['page']} | Pass | {p['finding']} |" for p in pages)
mdqa=f'''# Final paper export review

The IEEEtran PDF has 10 pages and the editable Word companion renders to 15 pages. Every page image was viewed individually at original detail. No material layout defect remains. The final PDF and Word use the same 15-reference numeric mapping. Word retains 73 native OMML expressions, seven numbered equation groups, four editable tables, three figures with captions, two amber notices, all M01–M16 fields, and six pending physical endpoints.

All 32 numeric cells in the main comparison table match the frozen paired-results JSON after declared rounding. Four participant medians, three RMSE differences/intervals and all four supplemental effects/intervals were checked against their frozen aggregates. Headline cohort, run and domain-failure counts were verified in both exports. Independent science review supports the underlying main and supplemental figures; this formatting pass did not rerun experiments.

Closed defects included missing Word figure captions, one-row table carryovers, a URL escape, raw mathematical typography and a PDF float overflow. The two result figures use verified F02 forecast comparison and F03 coverage/alarm figures from the filter-calibration lane. New Fig. 3 combines the whole conceptual rig image with a compact vector evidence chain, an explicit pending banner, separate command/report records and M01–M16 validation requirements. An editable SVG and PDF source are retained. The diagram contains no measured outcome or selected component claim. All three figures remain readable at publication size; the huge existing overview diagrams were not reduced into illegible labels.

The Fig. 3 update changed PDF pages 8–10 and Word pages 12–15; all seven were individually inspected. PDF pages 1–7 and Word pages 1–11 retain exact image SHA-256 identity with the previous individually inspected export. The final three-figure manuscript keeps all actual result numbers unchanged.

| Format | Page | Result | Observation |
|---|---:|---|---|
{rows}

Final manuscript SHA-256: `{sha(md)}`. PDF SHA-256: `{sha(pdf)}`. DOCX SHA-256: `{sha(word)}`. Exact per-page hashes, source hashes, reference mapping and numerical checks are in `FINAL_PAPER_QA.json`.

The editable Word layout was checked through the bundled LibreOffice renderer, not Microsoft Word. The PDF has normal IEEE float placement and unused final-reference space. Both amber PDF notices are intact on page 8; Fig. 3 and the pending physical-results table share page 9. No text clipping, missing glyphs or overlapping content was observed.

Scientific and release limits remain explicit: internally exposed retrospective data, upstream preprocessing limitations, synthetic parameter units, severe main Hovorka domain failures, post-primary exploratory fasting scope, original-family clamps, assumed stop authority and no measured hardware outcomes. Author, funding, ethics/governance, human AI-use review and final submission approval await the team. This receipt is not a clinical, metrology, authorship, originality or publication certification. Historical local overlap receipts apply to their recorded source versions and are not silently relabelled as checks of this final version.
'''
(QA/'FINAL_PAPER_QA.md').write_text(mdqa)
print(json.dumps({'status':'passed','pages':receipt['page_counts'],'outputs':receipt['outputs'][:3],'qa':str(QA)},indent=2))
